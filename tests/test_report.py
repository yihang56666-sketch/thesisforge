# -*- coding: utf-8 -*-
import _isolate  # noqa: F401

import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.report import build_report
from app.main import safe_report_filename
from app.report import json_params_short


class SafeReportFilenameTest(unittest.TestCase):
    def test_keeps_word_and_chinese_chars(self):
        self.assertEqual(safe_report_filename("基于ML的 研究:报告!? v1"), "基于ML的研究报告v1")

    def test_strips_windows_reserved_characters(self):
        self.assertEqual(safe_report_filename('题目 a/b\\c*d?"<>|'), "题目abcd")

    def test_empty_falls_back(self):
        self.assertEqual(safe_report_filename("///"), "报告")
        self.assertEqual(safe_report_filename("   "), "报告")

    def test_long_title_truncated(self):
        self.assertEqual(len(safe_report_filename("测" * 60)), 40)


class JsonParamsShortTest(unittest.TestCase):
    def test_truncates_to_four_params(self):
        out = json_params_short({"a": 1, "b": 2, "c": 3, "d": 4, "e": 5})
        self.assertEqual(out, "a=1, b=2, c=3, d=4")

    def test_empty_returns_default(self):
        self.assertEqual(json_params_short({}), "默认")
        self.assertEqual(json_params_short(None), "默认")


class ReportResearchDepthTest(unittest.TestCase):
    def _build(self, runs, dataset_meta):
        out = Path(tempfile.mkdtemp()) / "report.docx"
        build_report(out, "基于深度学习的分类研究", {}, dataset_meta, None, runs, {})
        doc = Document(str(out))
        return [p.text for p in doc.paragraphs]

    def test_without_ai_draft_includes_research_diagnostics(self):
        runs = [
            {
                "run_id": "base", "name": "基线", "group": "baseline",
                "config": {"task": "tabular_classification", "model": "logistic_regression",
                           "model_label": "逻辑回归", "group": "baseline", "params": {}},
                "summary": {"model_label": "逻辑回归", "group": "baseline",
                            "primary_metric": {"name": "accuracy", "value": 0.80},
                            "metrics": {"accuracy": 0.80}, "eval_source": "test"},
            },
            {
                "run_id": "imp", "name": "改进", "group": "improved",
                "config": {"task": "tabular_classification", "model": "mlp",
                           "model_label": "MLP", "group": "improved", "params": {}},
                "summary": {"model_label": "MLP", "group": "improved",
                            "primary_metric": {"name": "accuracy", "value": 0.86},
                            "metrics": {"accuracy": 0.86}, "eval_source": "test"},
            },
            {
                "run_id": "abl", "name": "消融", "group": "ablation",
                "config": {"task": "tabular_classification", "model": "mlp",
                           "model_label": "MLP", "group": "ablation", "params": {}},
                "summary": {"model_label": "MLP", "group": "ablation",
                            "primary_metric": {"name": "accuracy", "value": 0.82},
                            "metrics": {"accuracy": 0.82}, "eval_source": "test"},
            },
        ]
        dataset_meta = {
            "name": "示例分类集", "desc": "用于验证报告深度。", "columns": ["x1", "x2"],
            "target": "label", "stats": {"class_counts": {"yes": 80, "no": 20}},
        }
        texts = self._build(runs, dataset_meta)

        self.assertIn("4.2  综合对比与研究性解读", texts)
        self.assertTrue(any("0.8000" in t and "0.8600" in t and "7.5%" in t for t in texts))
        self.assertTrue(any("消融" in t and "0.8200" in t and "0.8600" in t for t in texts))
        self.assertTrue(any("类别不平衡" in t and "宏平均 F1" in t for t in texts))
        self.assertTrue(any("独立测试集" in t for t in texts))

    def test_report_flags_overfitting_from_epoch_history(self):
        runs = [
            {
                "run_id": "overfit", "name": "过拟合", "group": "baseline",
                "config": {"task": "image_classification", "model": "resnet18",
                           "model_label": "ResNet18", "group": "baseline", "params": {}},
                "summary": {"model_label": "ResNet18", "group": "baseline",
                            "primary_metric": {"name": "accuracy", "value": 0.86},
                            "metrics": {"accuracy": 0.86}, "eval_source": "test",
                            "epochs": [{"train_acc": 0.99, "val_acc": 0.80}]},
            }
        ]
        dataset_meta = {"name": "示例分类集", "columns": [], "target": "label"}
        texts = self._build(runs, dataset_meta)

        self.assertTrue(any("过拟合迹象" in t and "0.9900" in t and "0.8000" in t for t in texts))

    def test_report_includes_image_quality_limitations(self):
        dataset_meta = {
            "name": "示例图像集", "columns": [], "target": "label",
            "quality_warnings": [
                "有 3 张图像无法读取或损坏，训练前应删除或修复。",
                "有 2 张图像内容重复，建议在划分前去重，避免评估虚高。",
                "有 4 张图像尺寸异常，建议统一缩放或检查原始采集设置。",
            ],
        }
        texts = self._build([], dataset_meta)

        self.assertTrue(any("损坏图像" in t and "可用训练样本" in t for t in texts))
        self.assertTrue(any("内容重复图像" in t and "评估虚高" in t for t in texts))
        self.assertTrue(any("尺寸异常图像" in t and "缩放方案" in t for t in texts))

    def test_report_summarizes_repeated_experiments(self):
        runs = []
        for i, acc in enumerate([0.91, 0.93, 0.95], start=1):
            runs.append({
                "run_id": f"repeat-{i}", "name": f"改进-重复-{i}", "group": "improved",
                "config": {"task": "tabular_classification", "model": "mlp",
                           "model_label": "MLP", "group": "improved", "params": {},
                           "batch_id": "rep-1", "batch_kind": "repeats", "repeat_index": i},
                "summary": {"model_label": "MLP", "group": "improved",
                            "primary_metric": {"name": "accuracy", "value": acc},
                            "metrics": {"accuracy": acc}, "eval_source": "test"},
            })
        dataset_meta = {"name": "示例分类集", "columns": [], "target": "label"}
        texts = self._build(runs, dataset_meta)

        self.assertTrue(any("重复实验" in t and "0.930" in t and "0.020" in t for t in texts))

    def test_semantic_segmentation_report_describes_pixel_method_and_artifact(self):
        import io
        from PIL import Image

        run_dir = Path(tempfile.mkdtemp())
        img = Image.new("RGB", (8, 8), "white")
        img.save(run_dir / "segmentation_prediction.png")
        runs = [{
            "run_id": "seg", "name": "U-Net", "group": "baseline",
            "run_dir": str(run_dir),
            "config": {"task": "semantic_segmentation", "model": "unet",
                       "model_label": "U-Net 语义分割", "group": "baseline", "params": {}},
            "summary": {"model_label": "U-Net 语义分割", "group": "baseline",
                        "primary_metric": {"name": "iou", "value": 0.78},
                        "metrics": {"iou": 0.78, "dice": 0.86, "pixel_accuracy": 0.92},
                        "eval_source": "test"},
        }]
        out = Path(tempfile.mkdtemp()) / "report.docx"
        build_report(out, "基于深度学习的水面语义分割", {}, {"name": "示例分割集", "columns": [], "target": "mask"}, None, runs, {})
        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]

        self.assertTrue(any("像素级" in t and "IoU" in t and "Dice" in t for t in texts))
        self.assertTrue(any("U-Net" in t and "编码器" in t and "解码器" in t for t in texts))
        self.assertTrue(len(doc.inline_shapes) > 0)

    def test_report_generates_concrete_abstract_and_conclusion(self):
        runs = [
            {
                "run_id": "base", "name": "基线", "group": "baseline",
                "config": {"task": "tabular_classification", "model": "logistic_regression",
                           "model_label": "逻辑回归", "group": "baseline", "params": {}},
                "summary": {"model_label": "逻辑回归", "group": "baseline",
                            "primary_metric": {"name": "accuracy", "value": 0.80},
                            "metrics": {"accuracy": 0.80}, "eval_source": "test"},
            },
            {
                "run_id": "imp", "name": "改进", "group": "improved",
                "config": {"task": "tabular_classification", "model": "mlp",
                           "model_label": "MLP", "group": "improved", "params": {}},
                "summary": {"model_label": "MLP", "group": "improved",
                            "primary_metric": {"name": "accuracy", "value": 0.86},
                            "metrics": {"accuracy": 0.86}, "eval_source": "test"},
            },
        ]
        dataset_meta = {"name": "示例分类集", "desc": "用于验证报告质量。", "columns": ["x1", "x2"],
                        "target": "label", "stats": {"class_counts": {"yes": 80, "no": 20}}}
        out = Path(tempfile.mkdtemp()) / "report.docx"
        build_report(out, "基于深度学习的分类研究", {}, dataset_meta, None, runs, {})
        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]

        self.assertTrue(any("示例分类集" in t and "0.8000" in t and "0.8600" in t for t in texts))
        self.assertTrue(any("7.5%" in t and "逻辑回归" in t and "MLP" in t for t in texts))
        self.assertTrue(any("本文在示例分类集" in t and "0.8600" in t for t in texts))

    def test_report_cites_generated_analysis_figures(self):
        import io
        from PIL import Image

        run_dir = Path(tempfile.mkdtemp())
        Image.new("RGB", (8, 8), "white").save(run_dir / "curves.png")
        Image.new("RGB", (8, 8), "white").save(run_dir / "confusion_matrix.png")
        runs = [{
            "run_id": "run", "name": "实验", "group": "baseline", "run_dir": str(run_dir),
            "config": {"task": "image_classification", "model": "cnn",
                       "model_label": "CNN", "group": "baseline", "params": {}},
            "summary": {"model_label": "CNN", "group": "baseline",
                        "primary_metric": {"name": "accuracy", "value": 0.88},
                        "metrics": {"accuracy": 0.88}, "eval_source": "test"},
        }]
        out = Path(tempfile.mkdtemp()) / "report.docx"
        build_report(out, "基于图像分类的研究", {}, {"name": "示例图像集", "columns": [], "target": "label"}, None, runs, {})
        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]

        self.assertTrue(any("图 4-1 至 图 4-" in t and "CNN" in t for t in texts))

    def test_report_generates_task_aware_english_abstract_and_methodology(self):
        runs = [
            {
                "run_id": "base", "name": "基线", "group": "baseline",
                "config": {"task": "object_detection", "model": "yolov8n",
                           "model_label": "YOLOv8n", "group": "baseline", "params": {}},
                "summary": {"model_label": "YOLOv8n", "group": "baseline",
                            "primary_metric": {"name": "mAP50", "value": 0.82},
                            "metrics": {"mAP50": 0.82}, "eval_source": "test"},
            },
        ]
        dataset_meta = {"name": "示例航拍检测集", "columns": [], "target": "bbox"}
        out = Path(tempfile.mkdtemp()) / "report.docx"
        build_report(out, "Deep Learning Based Object Detection", {}, dataset_meta, None, runs, {})
        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]

        self.assertTrue(any("This thesis" in t and "0.8200" in t and "YOLOv8n" in t for t in texts))
        self.assertTrue(any("object detection" in t.lower() and "mAP50" in t for t in texts))
        self.assertTrue(any("bbox" in t and "mAP50" in t and "验证集" in t for t in texts))

    def test_report_generates_run_level_analysis_without_placeholder(self):
        runs = [
            {
                "run_id": "base", "name": "基线", "group": "baseline",
                "config": {"task": "tabular_classification", "model": "logistic_regression",
                           "model_label": "逻辑回归", "group": "baseline", "params": {}},
                "summary": {"model_label": "逻辑回归", "group": "baseline",
                            "primary_metric": {"name": "accuracy", "value": 0.80},
                            "metrics": {"accuracy": 0.80}, "eval_source": "test"},
            },
            {
                "run_id": "imp", "name": "改进", "group": "improved",
                "config": {"task": "tabular_classification", "model": "mlp",
                           "model_label": "MLP", "group": "improved", "params": {}},
                "summary": {"model_label": "MLP", "group": "improved",
                            "primary_metric": {"name": "accuracy", "value": 0.88},
                            "metrics": {"accuracy": 0.88, "macro_f1": 0.72},
                            "eval_source": "test",
                            "epochs": [{"train_acc": 0.98, "val_acc": 0.81}]},
            },
        ]
        dataset_meta = {"name": "示例分类集", "columns": ["x1", "x2"], "target": "label",
                        "stats": {"class_counts": {"yes": 80, "no": 20}}}
        texts = self._build(runs, dataset_meta)

        self.assertFalse(any("配置 AI 接口后，此处将自动生成实验解读" in t for t in texts))
        self.assertTrue(any("accuracy = 0.8800" in t and "macro_f1 = 0.7200" in t
                            and "过拟合" in t and "类别不平衡" in t for t in texts))

    def test_report_generates_structured_research_method_and_conclusion(self):
        runs = [
            {
                "run_id": "base", "name": "基线", "group": "baseline",
                "config": {"task": "tabular_classification", "model": "logistic_regression",
                           "model_label": "逻辑回归", "group": "baseline", "params": {}},
                "summary": {"model_label": "逻辑回归", "group": "baseline",
                            "primary_metric": {"name": "accuracy", "value": 0.80},
                            "metrics": {"accuracy": 0.80}, "eval_source": "test"},
            },
            {
                "run_id": "imp", "name": "改进", "group": "improved",
                "config": {"task": "tabular_classification", "model": "mlp",
                           "model_label": "MLP", "group": "improved", "params": {}},
                "summary": {"model_label": "MLP", "group": "improved",
                            "primary_metric": {"name": "accuracy", "value": 0.86},
                            "metrics": {"accuracy": 0.86}, "eval_source": "test"},
            },
        ]
        dataset_meta = {
            "name": "示例分类集", "columns": ["x1", "x2"], "target": "label", "n_rows": 200,
            "stats": {"class_counts": {"yes": 80, "no": 20}},
        }
        texts = self._build(runs, dataset_meta)

        self.assertIn("1.3  研究方法与技术路线", texts)
        self.assertIn("4.2  综合对比与研究性解读", texts)
        self.assertIn("5.1  工作总结", texts)
        self.assertIn("5.2  主要结论", texts)
        self.assertIn("5.3  研究局限", texts)
        self.assertIn("5.4  未来展望", texts)
        self.assertTrue(any("数据获取与质量控制" in t and "对比与消融验证" in t for t in texts))
        self.assertTrue(any("验证集用于模型与超参数选择" in t and "测试集只用于最终评估" in t for t in texts))
        self.assertTrue(any("样本规模约为 200 条" in t for t in texts))
        self.assertTrue(any("部署延迟" in t and "能耗" in t for t in texts))

    def test_report_generates_task_specific_literature_review(self):
        runs = [
            {
                "run_id": "base", "name": "目标检测基线", "group": "baseline",
                "config": {"task": "object_detection", "model": "yolov8n",
                           "model_label": "YOLOv8n", "group": "baseline", "params": {}},
                "summary": {"model_label": "YOLOv8n", "group": "baseline",
                            "primary_metric": {"name": "mAP50", "value": 0.82},
                            "metrics": {"mAP50": 0.82}, "eval_source": "test"},
            },
            {
                "run_id": "imp", "name": "分割对照", "group": "improved",
                "config": {"task": "semantic_segmentation", "model": "unet",
                           "model_label": "U-Net 语义分割", "group": "improved", "params": {}},
                "summary": {"model_label": "U-Net 语义分割", "group": "improved",
                            "primary_metric": {"name": "iou", "value": 0.78},
                            "metrics": {"iou": 0.78, "dice": 0.86, "pixel_accuracy": 0.92},
                            "eval_source": "test"},
            },
        ]
        dataset_meta = {"name": "示例视觉数据集", "columns": [], "target": "bbox"}
        texts = self._build(runs, dataset_meta)

        self.assertTrue(any("文献综述" in t for t in texts))
        self.assertTrue(any("目标检测" in t and "YOLO" in t and "端到端" in t for t in texts))
        self.assertTrue(any("语义分割" in t and "U-Net" in t and "IoU" in t for t in texts))

    def test_report_summarizes_dataset_quality_warnings(self):
        runs = [{
            "run_id": "base", "name": "基线", "group": "baseline",
            "config": {"task": "tabular_classification", "model": "logistic_regression",
                       "model_label": "逻辑回归", "group": "baseline", "params": {}},
            "summary": {"model_label": "逻辑回归", "group": "baseline",
                        "primary_metric": {"name": "accuracy", "value": 0.80},
                        "metrics": {"accuracy": 0.80}, "eval_source": "test"},
        }]
        dataset_meta = {
            "name": "示例数据集", "columns": ["feature"], "target": "label", "n_rows": 150,
            "quality_warnings": [
                "存在 3 个缺失值，需在预处理阶段明确填充或删除策略。",
                "存在 2 条重复样本，建议在划分前去重，避免训练/评估信息泄漏。",
            ],
        }
        texts = self._build(runs, dataset_meta)

        self.assertTrue(any("数据质量检查" in t and "缺失值" in t and "重复样本" in t for t in texts))


if __name__ == "__main__":
    unittest.main()
